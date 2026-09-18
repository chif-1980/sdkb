import io
import json
from types import SimpleNamespace

import pytest
from docx import Document

from yuxi.product_chat.meeting_export import export_docx
from yuxi.product_chat.meeting_service import (
    build_followup,
    call_cited_json,
    call_json,
    transcript_chunks,
    validate_citations,
)
from yuxi.product_chat.meeting_sources import (
    MeetingSourceError,
    PublicResolver,
    parse_buddy,
    parse_tingwu,
    read_attachment,
    read_text,
    validate_url,
)
from yuxi.product_chat.meeting_web_reader import merge_visible


@pytest.mark.parametrize(
    "url",
    [
        "http://127.0.0.1/test",
        "http://192.168.6.95/",
        "http://169.254.169.254/",
        "http://[::1]/",
        "http://localhost/",
        "http://service.internal/",
        "file:///etc/passwd",
        "https://user:password@example.com/",
        "https://example.com:9000/",
    ],
)
def test_rejects_non_public_addresses(url):
    with pytest.raises(MeetingSourceError, match="地址|端口|凭据"):
        validate_url(url)


async def test_dns_mixed_public_private_answers_are_rejected(monkeypatch):
    async def resolve(*args, **kwargs):
        return [{"host": "93.184.216.34"}, {"host": "10.0.0.1"}]

    monkeypatch.setattr("aiohttp.resolver.DefaultResolver.resolve", resolve)
    with pytest.raises(MeetingSourceError) as failure:
        await PublicResolver().resolve("meeting.example.com")
    assert failure.value.code == "UNSAFE_ADDRESS"


def test_new_public_platform_does_not_require_domain_adapter():
    assert validate_url("https://new-meeting.example.com/share/abc")


def test_followup_keeps_explicit_actions_separate_from_knowledge_suggestions():
    result = build_followup(
        {
            "actionItems": [
                {
                    "title": "补充测试方案",
                    "assigneeName": "张工",
                    "dueDate": "未明确",
                    "evidence": "[S1-P2]",
                }
            ],
            "knowledgeSuggestions": [
                {"title": "更新部署限制", "reason": "会议提出了新的限制条件", "evidence": "[S1-P3]"}
            ],
        },
        coordinator_id=7,
        coordinator_name="会议上传者",
    )

    assert result["coordinator"] == {"userId": "7", "displayName": "会议上传者"}
    assert result["tasks"] == [
        {
            "id": "task-1",
            "title": "补充测试方案",
            "assignee": None,
            "assigneeSuggestion": "张工",
            "dueDate": None,
            "status": "OPEN",
            "sourceRefs": ["S1-P2"],
        }
    ]
    assert result["knowledgeSuggestions"][0]["status"] == "PENDING_MAINTAINER"


def test_buddy_preserves_every_segment_and_real_locators():
    rows = [
        {"message": f"第 {i} 段", "speaker": "甲", "startTime": i * 1000, "endTime": (i + 1) * 1000}
        for i in range(1280)
    ]
    source = parse_buddy(
        {"success": True, "result": {"messageList": rows, "summary": "平台总结"}}, "https://example.com"
    )
    assert len(source["paragraphs"]) == 1280
    assert source["paragraphs"][-1]["id"] == "P1280"
    assert source["paragraphs"][-1]["startMs"] == 1279000
    assert source["platformSummary"] == "平台总结"


@pytest.mark.parametrize(
    "rows,summary,code",
    [([], "", "EMPTY"), ([], "总结", "SUMMARY_ONLY"), ([{"message": "正文"}, {"message": ""}], "", "PARTIAL")],
)
def test_incomplete_sources_never_claim_success(rows, summary, code):
    with pytest.raises(MeetingSourceError) as failure:
        parse_buddy({"success": True, "result": {"messageList": rows, "summary": summary}}, "https://example.com")
    assert failure.value.code == code


def test_tingwu_rejects_truncated_duration():
    payload = {
        "code": "0",
        "data": {
            "duration": 10067,
            "result": json.dumps({"pg": [{"sc": [{"tc": "只读取了第一分钟", "bt": 0, "et": 60000}]}]}),
        },
    }
    with pytest.raises(MeetingSourceError) as failure:
        parse_tingwu(payload, "https://example.com")
    assert failure.value.code == "PARTIAL"


def test_chunking_covers_long_unbroken_text_without_losing_locator():
    source = read_text("长" * 19000 + "\n最后一段")
    chunks = transcript_chunks([source])
    assert "".join(p["text"] for chunk in chunks for p in chunk) == "长" * 19000 + "最后一段"
    assert all(sum(len(p["text"]) for p in chunk) <= 7000 for chunk in chunks)
    assert chunks[2][0]["id"] == "S1-P1"
    assert chunks[-1][-1]["id"] == "S1-P2"
    with pytest.raises(ValueError):
        validate_citations("虚假引用 [S1-P99]", {"S1-P1"})


def test_overlapping_pages_keep_repeated_utterances():
    def ps(*texts):
        return [{"text": text} for text in texts]

    assert merge_visible(ps("甲", "同意", "乙"), ps("乙", "同意", "结束")) == ps("甲", "同意", "乙", "同意", "结束")
    assert merge_visible(ps("甲", "同意", "乙"), ps("同意", "乙")) == ps("甲", "同意", "乙")


def test_history_references_cannot_claim_current_or_unselected_evidence():
    validate_citations("本次未决定 [S1-P1]；历史讨论 [H1]", {"S1-P1", "H1"})
    for body in ("历史原来的段落号 [S1-P2]", "未选会议 [H2]"):
        with pytest.raises(ValueError):
            validate_citations(body, {"S1-P1", "H1"})


async def test_invalid_reference_is_rechecked_once_against_same_evidence():
    class Model:
        calls = []

        async def call(self, messages):
            self.calls.append(json.loads(messages[1]["content"]))
            return SimpleNamespace(
                content=json.dumps({"body": "待确认 [S1-P2]" if len(self.calls) == 1 else "待确认 [S1-P1]"})
            )

    model = Model()
    evidence = {"transcript": [{"id": "S1-P1", "text": "仅为建议"}]}
    result = await call_cited_json(model, "整理纪要", evidence, {"S1-P1"})
    assert result["body"] == "待确认 [S1-P1]"
    assert len(model.calls) == 2
    assert model.calls[1]["transcript"] == evidence["transcript"]
    assert "S1-P2" in model.calls[1]["validationError"]


def test_docx_input_keeps_paragraph_table_order_and_export_saved_body(tmp_path):
    doc = Document()
    doc.add_paragraph("首先讨论交付")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "负责人", "待确认"
    doc.add_paragraph("最后确认下周再讨论")
    path = tmp_path / "meeting.docx"
    doc.save(path)
    source = read_attachment(str(path), path.name)
    assert [p["text"] for p in source["paragraphs"]] == ["首先讨论交付", "负责人 | 待确认", "最后确认下周再讨论"]
    result = {
        "title": "编辑后的会议",
        "meetingType": "内部管理",
        "body": "## 行动清单\n\n| 事项 | 负责人 |\n| --- | --- |\n| 再讨论 [S1-P3] | 待确认 |",
    }
    exported = Document(io.BytesIO(export_docx(result, [source], 2)))
    assert exported.tables[0].cell(1, 1).text == "待确认"
    text = "\n".join(p.text for p in exported.paragraphs)
    assert "编辑后的会议" in text and "版本 2" in text
    assert "[S1-P3] 未提供 | P3" in text and "最后确认下周再讨论" in text


@pytest.mark.parametrize("invalid", ['{"body":"unfinished', '{"body":"raw\nnewline"}', "[]", ""])
async def test_malformed_json_retries_once_with_original_evidence(invalid):
    from unittest.mock import AsyncMock

    model = SimpleNamespace(
        call=AsyncMock(
            side_effect=[
                SimpleNamespace(content=invalid),
                SimpleNamespace(content='{"body":"依据 [S1-P1]"}'),
            ]
        )
    )
    evidence = {"transcript": [{"id": "S1-P1", "text": "只是建议"}]}
    result = await call_cited_json(model, "整理纪要", evidence, {"S1-P1"})
    assert result["body"] == "依据 [S1-P1]"
    assert model.call.await_count == 2
    for call in model.call.call_args_list:
        assert json.loads(call.args[0][1]["content"])["transcript"] == evidence["transcript"]


async def test_repeated_invalid_json_has_specific_failure_and_bounded_retry():
    from unittest.mock import AsyncMock

    model = SimpleNamespace(call=AsyncMock(return_value=SimpleNamespace(content='{"body":')))
    with pytest.raises(ValueError) as failure:
        await call_json(model, "整理纪要", {})
    assert failure.value.code == "MODEL_OUTPUT_INVALID"
    assert "无需更换链接" in str(failure.value)
    assert model.call.await_count == 2


async def test_truncated_output_is_not_accepted_even_if_json_parses():
    from unittest.mock import AsyncMock

    model = SimpleNamespace(
        call=AsyncMock(
            return_value=SimpleNamespace(
                content='{"body":"partial"}',
                metadata={"finish_reason": "length"},
            )
        )
    )
    with pytest.raises(ValueError) as failure:
        await call_json(model, "整理纪要", {})
    assert failure.value.code == "MODEL_OUTPUT_TRUNCATED"
    assert model.call.await_count == 2


async def test_model_connection_failure_is_not_misreported_as_json_error():
    from unittest.mock import AsyncMock

    model = SimpleNamespace(call=AsyncMock(side_effect=TimeoutError("upstream timeout")))
    with pytest.raises(TimeoutError):
        await call_json(model, "整理纪要", {})
    assert model.call.await_count == 1


@pytest.mark.parametrize(
    "state,changed,resume",
    [
        ("failed", False, True),
        ("cancelled", False, True),
        ("completed", False, False),
        ("failed", True, False),
    ],
)
async def test_retry_reuses_notes_only_for_unchanged_incomplete_meeting(monkeypatch, state, changed, resume):
    from unittest.mock import AsyncMock, Mock
    from yuxi.product_chat.meeting_repository import MeetingRepository, ProductChatRepository
    from yuxi.product_chat.schemas import SendMessageRequest

    db = SimpleNamespace(
        scalar=AsyncMock(return_value=None), execute=AsyncMock(), flush=AsyncMock(), add_all=Mock(), add=Mock()
    )
    conversation = SimpleNamespace(id=1, status="ACTIVE", title="会议")
    monkeypatch.setattr(ProductChatRepository, "require_conversation", AsyncMock(return_value=conversation))
    notes = [{"summary": "只是提议 [S1-P1]"}]
    parent = SimpleNamespace(
        id="parent",
        conversation_id="conversation",
        state=state,
        sources=[{"title": "已读取资料"}],
        result=None,
        input={"content": "原始请求", "chunkNotes": notes},
    )
    repo = MeetingRepository(db)
    monkeypatch.setattr(repo, "require", AsyncMock(return_value=parent))
    record = await repo.create(
        "conversation",
        SimpleNamespace(id=1),
        SendMessageRequest(
            content="修改后的请求" if changed else "原始请求",
            meetingId="parent",
            skillId="MEETING_ANALYSIS",
        ),
    )
    assert record.input["chunkNotes"] == (notes if resume else [])
    assert parent.input["chunkNotes"] == notes


async def test_chat_adapter_preserves_finish_reason_for_truncation_check():
    from unittest.mock import AsyncMock
    from yuxi.models.chat import LangChainChatAdapter

    upstream = SimpleNamespace(
        ainvoke=AsyncMock(
            return_value=SimpleNamespace(
                text="partial",
                response_metadata={"finish_reason": "length"},
            )
        )
    )
    result = await LangChainChatAdapter(upstream, model_name="test").call("hello")
    assert result.metadata["finish_reason"] == "length"


async def test_worker_persists_specific_format_failure_and_last_step(monkeypatch):
    from contextlib import asynccontextmanager
    from unittest.mock import AsyncMock
    from yuxi.product_chat import meeting_service as service

    record = SimpleNamespace(
        state="running", progress={"message": "核对决定", "updatedAt": "old"}, error=None, message_id="msg"
    )
    message = SimpleNamespace(content="")
    db = SimpleNamespace(get=AsyncMock(return_value=record), scalar=AsyncMock(return_value=message))

    @asynccontextmanager
    async def session():
        yield db

    monkeypatch.setattr(service.pg_manager, "get_async_session_context", session)
    monkeypatch.setattr(service, "_analyze", AsyncMock(side_effect=service.MeetingOutputError()))
    monkeypatch.setattr(service, "has_cancel_signal", AsyncMock(return_value=False))
    event = AsyncMock()
    monkeypatch.setattr(service, "append_run_stream_event", event)
    await service.process_meeting_run({}, "meeting")
    assert record.state == "failed"
    assert record.error["code"] == "MODEL_OUTPUT_INVALID"
    assert record.progress["failedStep"] == "核对决定"
    assert record.progress["updatedAt"] != "old"
    assert message.content == record.error["message"]
    assert event.call_args.args[1] == "failed"
