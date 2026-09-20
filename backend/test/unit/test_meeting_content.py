import io

from docx import Document

from yuxi.product_chat.meeting_content import meeting_body, split_task_section
from yuxi.product_chat.meeting_export import export_docx
from yuxi.product_chat.meeting_management import merge_followups
from yuxi.product_chat.meeting_service import build_followup


def test_tasks_replace_stale_tables_without_removing_decisions_or_code():
    body = (
        "## 明确决定\n原决定\n```md\n## 行动清单\n示例代码\n```\n\n"
        "## 三、行动清单\n旧负责人 旧日期\n### 细节\n旧细节\n"
        "## 待确认问题\n人数待确认\n## 待办事项\n重复表\n## 业务分析\n分析保留"
    )
    result = {
        "title": "会议",
        "body": body,
        "followup": {
            "tasks": [
                {
                    "id": "task-1",
                    "title": "更新|手册",
                    "content": "核对\n附录",
                    "assignee": {"displayName": "新负责人"},
                    "dueDate": "2026-09-21",
                    "status": "DONE",
                    "reviewStatus": "CONFIRMED",
                    "sourceRefs": ["S1-P1"],
                }
            ]
        },
    }
    rendered = meeting_body(result)
    assert "旧负责人" not in rendered and "旧细节" not in rendered and "重复表" not in rendered
    assert "原决定" in rendered and "示例代码" in rendered and "人数待确认" in rendered and "分析保留" in rendered
    assert "新负责人" in rendered and "2026-09-21" in rendered and "已确认 / 已完成" in rendered
    assert rendered.count("## 待办事项") == 1
    assert meeting_body({**result, "body": rendered}) == rendered
    exported = Document(
        io.BytesIO(
            export_docx(
                result,
                [
                    {
                        "title": "转写",
                        "platform": "文字",
                        "paragraphs": [{"id": "P1", "speaker": "发言人", "text": "核对附录"}],
                    }
                ],
                3,
            )
        )
    )
    assert len(exported.tables) == 1
    cells = " ".join(c.text for row in exported.tables[0].rows for c in row.cells)
    assert "更新|手册" in cells and "新负责人" in cells and "已完成" in cells
    assert "[S1-P1]" in "\n".join(p.text for p in exported.paragraphs)


def test_legacy_without_structured_tasks_and_empty_tasks_are_distinct():
    original = "## 行动清单\n历史唯一内容"
    assert meeting_body({"body": original}) == original
    assert "暂无待办事项" in meeting_body({"body": original, "followup": {"tasks": []}})
    assert split_task_section("正文\n## 待确认问题\n未确认") == ("正文", "## 待确认问题\n未确认")


def test_ai_revision_is_a_proposal_preserving_identity_human_changes_and_delivery():
    old = {
        "id": "stable-1",
        "title": "人工修订的任务",
        "status": "DONE",
        "reviewStatus": "CONFIRMED",
        "dueDate": "2026-09-21",
        "assignee": {"displayName": "甲"},
        "delivery": {"feishuTaskId": "remote-1"},
    }
    incoming = build_followup(
        {
            "actionItems": [
                {
                    "taskId": "stable-1",
                    "title": "修订后任务",
                    "content": "补充要求",
                    "assigneeName": "乙",
                    "dueDate": "2026-09-23",
                    "evidence": "[S1-P1]",
                }
            ]
        },
        coordinator_id=1,
        coordinator_name="甲",
    )
    merged = merge_followups([old], incoming["tasks"], run_id="MT-new")
    assert len(merged) == 1
    assert {k: merged[0][k] for k in old} == old
    assert merged[0]["aiProposal"]["title"] == "修订后任务"
    assert merged[0]["aiProposal"]["sourceMeetingId"] == "MT-new"
    assert "aiProposal" not in old


async def test_task_verification_uses_original_evidence_and_preserves_target_ids(monkeypatch):
    from yuxi.product_chat import meeting_service

    async def review(model, prompt, data, allowed):
        assert data["originalEvidence"]["S1-P1"]["text"] == "建议后续讨论，人员与日期待定"
        assert allowed == {"S1-P1"}
        return {
            "actionItems": [
                {
                    "title": "后续讨论",
                    "content": "建议，待确认",
                    "assigneeName": "待确认",
                    "dueDate": "待确认",
                    "evidence": "[S1-P1]",
                    "taskId": "model-made-up-id",
                }
            ]
        }

    monkeypatch.setattr(meeting_service, "call_cited_json", review)
    checked = await meeting_service.verify_action_items(
        None,
        [{"title": "后续讨论", "taskId": "stable-1", "evidence": "[S1-P1]"}],
        {"S1-P1": {"text": "建议后续讨论，人员与日期待定"}},
        {},
        "整理会议",
    )
    assert checked[0]["taskId"] == "stable-1"
    assert checked[0]["dueDate"] == "待确认"


async def test_export_retained_task_refs_resolve_to_original_run(monkeypatch):
    from types import SimpleNamespace
    from yuxi.product_chat.meeting_repository import MeetingRepository

    def source(text):
        return {"title": text, "platform": "文字", "paragraphs": [{"id": "P1", "text": text, "speaker": "原发言人"}]}

    record = SimpleNamespace(
        id="new",
        conversation_id="conversation",
        sources=[source("新原文")],
        result={
            "body": "正文依据 [S1-P1]",
            "followup": {
                "tasks": [{"id": "old-task", "title": "保留的待办", "sourceMeetingId": "old", "sourceRefs": ["S1-P1"]}]
            },
        },
    )
    repo = MeetingRepository(None)

    async def require(source_id, user_id):
        assert source_id == "old" and user_id == 7
        return SimpleNamespace(conversation_id="conversation", sources=[source("真正的旧原文")])

    monkeypatch.setattr(repo, "require", require)
    result, sources = await repo.export_materials(record, 7)
    assert result["followup"]["tasks"][0]["sourceRefs"] == ["S2-P1"]
    assert record.result["followup"]["tasks"][0]["sourceRefs"] == ["S1-P1"]
    assert sources[1]["paragraphs"][0]["text"] == "真正的旧原文"
    doc = Document(io.BytesIO(export_docx(result, sources, 2)))
    assert any("[S2-P1]" in p.text and "真正的旧原文" in p.text for p in doc.paragraphs)
