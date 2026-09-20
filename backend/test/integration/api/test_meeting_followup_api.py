"""Live HTTP contract checks using disposable data; no model jobs or external messages."""

import os
from uuid import uuid4

import httpx
import pytest
from sqlalchemy import delete

from yuxi.storage.postgres.manager import pg_manager
from yuxi.storage.postgres.models_business import User, Department
from yuxi.storage.postgres.models_product import MeetingRecord, MeetingRevision, ProductConversation, ProductMessage
from yuxi.utils.auth_utils import AuthUtils

pytestmark = pytest.mark.integration


@pytest.fixture
async def followup_clients():
    pg_manager.initialize()
    prefix = f"followup-test-{uuid4().hex[:10]}"
    async with pg_manager.get_async_session_context() as db:
        department = Department(name=prefix)
        db.add(department)
        await db.flush()
        users = [
            User(
                username=f"{prefix}-{i}", uid=f"{prefix}-{i}", password_hash="not-a-login", department_id=department.id
            )
            for i in range(2)
        ]
        db.add_all(users)
        await db.flush()
        ids, department_id = [u.id for u in users], department.id
        conversation = ProductConversation(owner_user_id=ids[0], title=prefix)
        db.add(conversation)
        await db.flush()
        cid = conversation.conversation_id
        question = ProductMessage(conversation_id=cid, role="USER", content="测试会议")
        answer = ProductMessage(conversation_id=cid, role="ASSISTANT", content="测试结果", answer_status="SUPPORTED")
        db.add_all([question, answer])
        await db.flush()
        mid = f"MT-{uuid4().hex}"
        db.add(
            MeetingRecord(
                id=mid,
                conversation_id=cid,
                message_id=answer.message_id,
                user_message_id=question.message_id,
                request_id=prefix,
                state="completed",
                version=1,
                progress={"message": "完成"},
                sources=[],
                input={},
                result={
                    "title": prefix,
                    "body": "测试结果",
                    "followup": {
                        "coordinator": {"userId": str(ids[0]), "displayName": prefix},
                        "tasks": [
                            {
                                "id": "task-1",
                                "title": "核对资料",
                                "dueDate": None,
                                "dueDateSuggestion": "2026年9月21日前",
                                "assignee": None,
                                "status": "OPEN",
                                "sourceRefs": [],
                            }
                        ],
                        "knowledgeSuggestions": [],
                    },
                },
            )
        )
    clients = [
        httpx.AsyncClient(
            base_url=os.getenv("TEST_BASE_URL", "http://localhost:5050"),
            timeout=30,
            cookies={
                "enterprise_assistant_session": AuthUtils.create_access_token(
                    {"sub": str(uid), "token_kind": "enterprise_assistant"}
                )
            },
        )
        for uid in ids
    ]
    try:
        yield clients, mid
    finally:
        for client in clients:
            await client.aclose()
        async with pg_manager.get_async_session_context() as db:
            await db.execute(delete(MeetingRevision).where(MeetingRevision.meeting_id == mid))
            await db.execute(delete(MeetingRecord).where(MeetingRecord.id == mid))
            await db.execute(delete(ProductMessage).where(ProductMessage.conversation_id == cid))
            await db.execute(delete(ProductConversation).where(ProductConversation.conversation_id == cid))
            await db.execute(delete(User).where(User.id.in_(ids)))
            await db.execute(delete(Department).where(Department.id == department_id))
        await pg_manager.close()


async def test_activity_is_owned_dates_validated_and_revision_saved(followup_clients):
    (owner, stranger), mid = followup_clients
    activity = await owner.get("/api/chat/meeting-activity")
    assert activity.status_code == 200, activity.text
    assert [task["id"] for task in activity.json()["tasks"]] == [mid]
    assert (await stranger.get("/api/chat/meeting-activity")).json()["tasks"] == []
    assert (await stranger.get(f"/api/chat/meetings/{mid}/followup-directory")).status_code == 404
    assert (await owner.get(f"/api/chat/meetings/{mid}/followup-directory")).status_code == 403
    patch = {"version": 1, "tasks": [{"id": "task-1", "title": "核对资料", "dueDate": "2026-02-30"}]}
    assert (await owner.patch(f"/api/chat/meetings/{mid}/followup", json=patch)).status_code == 422
    patch["tasks"][0]["dueDate"] = "2026-09-30"
    assert (await stranger.patch(f"/api/chat/meetings/{mid}/followup", json=patch)).status_code == 404
    saved = await owner.patch(f"/api/chat/meetings/{mid}/followup", json=patch)
    assert saved.status_code == 200, saved.text
    assert saved.json()["meeting"]["version"] == 2
    assert saved.json()["meeting"]["result"]["followup"]["tasks"][0]["dueDate"] == "2026-09-30"
    assert (await owner.patch(f"/api/chat/meetings/{mid}/followup", json=patch)).status_code == 409
    patch["version"] = 2
    patch["tasks"][0]["dueDate"] = None
    cleared = await owner.patch(f"/api/chat/meetings/{mid}/followup", json=patch)
    assert cleared.status_code == 200, cleared.text
    reloaded = await owner.get(f"/api/chat/meetings/{mid}")
    task = reloaded.json()["meeting"]["result"]["followup"]["tasks"][0]
    assert task["dueDate"] is None
    assert task["dueDateEdited"] is True
    assert task["dueDateSuggestion"] == "2026年9月21日前"


async def test_unified_minutes_copy_export_edit_and_ai_proposal_do_not_send(followup_clients):
    import io
    from docx import Document
    from sqlalchemy import select

    from yuxi.storage.postgres.models_product import ProductMessage

    (owner, stranger), mid = followup_clients
    async with pg_manager.get_async_session_context() as db:
        record = await db.get(MeetingRecord, mid)
        record.result = {
            **record.result,
            "body": "## 明确决定\n决定保留\n## 行动清单\n错误旧任务\n## 待确认问题\n人数待确认",
        }
    patch = {
        "version": 1,
        "tasks": [{"id": "task-1", "title": "最新待办", "dueDate": "2026-09-22", "status": "IN_PROGRESS"}],
    }
    saved = await owner.patch(f"/api/chat/meetings/{mid}/followup", json=patch)
    assert saved.status_code == 200, saved.text
    result = saved.json()["meeting"]["result"]
    assert "错误旧任务" not in result["body"] and "最新待办" in result["body"]
    assert "2026-09-22" in result["body"]
    assert (await stranger.get(f"/api/chat/meetings/{mid}/markdown?version=2")).status_code == 404
    assert (await owner.get(f"/api/chat/meetings/{mid}/markdown?version=1")).status_code == 409
    copied = await owner.get(f"/api/chat/meetings/{mid}/markdown?version=2")
    assert copied.json()["body"] == result["body"]
    exported = await owner.get(f"/api/chat/meetings/{mid}/export?version=2")
    assert exported.status_code == 200
    doc = Document(io.BytesIO(exported.content))
    assert len(doc.tables) == 1
    assert "最新待办" in " ".join(c.text for r in doc.tables[0].rows for c in r.cells)
    edited = await owner.patch(
        f"/api/chat/meetings/{mid}",
        json={
            "version": 2,
            "title": "手动修改标题",
            "meetingType": "内部管理",
            "body": "## 明确决定\n新的正文\n## 待办事项\n试图写入另一份任务\n## 待确认问题\n待确认保留",
        },
    )
    assert edited.status_code == 200, edited.text
    assert "试图写入另一份任务" not in edited.json()["meeting"]["result"]["body"]
    assert "最新待办" in edited.json()["meeting"]["result"]["body"]
    async with pg_manager.get_async_session_context() as db:
        record = await db.get(MeetingRecord, mid)
        task = {
            **record.result["followup"]["tasks"][0],
            "aiProposal": {
                "title": "建议修改标题",
                "content": "建议补充",
                "assigneeSuggestion": "待核对成员",
                "dueDate": "2026-09-23",
                "sourceRefs": [],
                "sourceMeetingId": mid,
            },
        }
        record.result = {**record.result, "followup": {**record.result["followup"], "tasks": [task]}}
    applied = await owner.patch(
        f"/api/chat/meetings/{mid}/followup", json={**patch, "version": 3, "action": "APPLY_AI", "taskId": "task-1"}
    )
    assert applied.status_code == 200, applied.text
    task = applied.json()["meeting"]["result"]["followup"]["tasks"][0]
    assert task["title"] == "建议修改标题" and task["id"] == "task-1"
    assert task["assignee"] is None and task["reviewStatus"] == "PENDING"
    assert task["dueDate"] == "2026-09-23" and not task["delivery"]["feishuTaskId"]
    assert "aiProposal" not in task
    async with pg_manager.get_async_session_context() as db:
        record = await db.get(MeetingRecord, mid)
        message = await db.scalar(select(ProductMessage).where(ProductMessage.message_id == record.message_id))
        assert "建议修改标题" in message.content and "最新待办" not in message.content


async def test_ai_proposal_keeps_existing_remote_task_and_can_be_dismissed(followup_clients):
    (owner, _), mid = followup_clients
    proposal = {
        "title": "修改执行要求",
        "content": "核对后交付",
        "assigneeSuggestion": "新负责人",
        "dueDate": "2026-09-25",
        "sourceRefs": [],
        "sourceMeetingId": mid,
    }
    async with pg_manager.get_async_session_context() as db:
        record = await db.get(MeetingRecord, mid)
        task = {
            **record.result["followup"]["tasks"][0],
            "reviewStatus": "CONFIRMED",
            "status": "DONE",
            "delivery": {"feishuTaskId": "fake-never-call-feishu", "notification": "SENT"},
            "aiProposal": proposal,
        }
        record.result = {**record.result, "followup": {**record.result["followup"], "tasks": [task]}}
    patch = {
        "version": 1,
        "action": "APPLY_AI",
        "taskId": "task-1",
        "tasks": [{"id": "task-1", "title": "核对资料", "status": "DONE"}],
    }
    applied = await owner.patch(f"/api/chat/meetings/{mid}/followup", json=patch)
    assert applied.status_code == 200, applied.text
    task = applied.json()["meeting"]["result"]["followup"]["tasks"][0]
    assert task["id"] == "task-1" and task["status"] == "DONE" and task["reviewStatus"] == "CONFIRMED"
    assert task["delivery"]["feishuTaskId"] == "fake-never-call-feishu"
    assert task["delivery"]["pendingUpdate"] is True and task["delivery"]["syncStatus"] == "PENDING"
    async with pg_manager.get_async_session_context() as db:
        record = await db.get(MeetingRecord, mid)
        record.result = {
            **record.result,
            "followup": {
                **record.result["followup"],
                "tasks": [{**task, "aiProposal": {**proposal, "title": "应被拒绝的建议"}}],
            },
        }
    patch.update(
        {
            "version": 2,
            "action": "DISMISS_AI",
            "tasks": [
                {
                    "id": "task-1",
                    "title": task["title"],
                    "content": task["content"],
                    "dueDate": task["dueDate"],
                    "status": "DONE",
                }
            ],
        }
    )
    dismissed = await owner.patch(f"/api/chat/meetings/{mid}/followup", json=patch)
    assert dismissed.status_code == 200, dismissed.text
    saved = dismissed.json()["meeting"]["result"]["followup"]["tasks"][0]
    assert saved["title"] == "修改执行要求" and "aiProposal" not in saved
    assert saved["delivery"] == task["delivery"]
    patch["version"] = 3
    assert (await owner.patch(f"/api/chat/meetings/{mid}/followup", json=patch)).status_code == 409
