"""Export the saved meeting version, including cited original evidence."""

import io
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt


def export_docx(result: dict, sources: list[dict], version: int) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    document.styles["Normal"].font.size = Pt(10.5)
    document.styles["Normal"].paragraph_format.space_after = Pt(6)
    document.add_heading(result.get("title") or "会议纪要", 0)
    document.add_paragraph(f"版本 {version}  |  {result.get('meetingType', '未提供')}")
    parser = MarkdownIt("commonmark").enable("table")
    tokens = parser.parse(result["body"])
    table, row, cell = None, None, None
    heading = 0
    paragraph = None
    for token in tokens:
        if token.type == "heading_open":
            heading = min(int(token.tag[1:]), 3)
        elif token.type == "table_open":
            table = document.add_table(rows=0, cols=0)
            table.style = "Table Grid"
        elif token.type == "tr_open" and table is not None:
            row = table.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            cell = 0
        elif token.type in ("th_open", "td_open") and table is not None:
            while len(table.columns) <= cell:
                table.add_column(Cm(3))
            paragraph = row.cells[cell].paragraphs[0]
            cell += 1
            if token.type == "th_open":
                repeat = OxmlElement("w:tblHeader")
                row._tr.get_or_add_trPr().append(repeat)
        elif token.type == "table_close":
            if table is not None and table.columns:
                width = Cm(16.6 / len(table.columns))
                table.autofit = False
                for column in table.columns:
                    column.width = width
                for table_row in table.rows:
                    for table_cell in table_row.cells:
                        table_cell.width = width
            table, row, cell, paragraph = None, None, None, None
        elif token.type == "paragraph_open" and table is None:
            paragraph = document.add_paragraph()
        elif token.type == "inline":
            if heading:
                paragraph = document.add_heading(level=heading)
                heading = 0
            if paragraph is None:
                paragraph = document.add_paragraph()
            bold = False
            for child in token.children or []:
                if child.type == "strong_open":
                    bold = True
                elif child.type == "strong_close":
                    bold = False
                elif child.type in ("text", "code_inline"):
                    paragraph.add_run(child.content).bold = bold
                elif child.type in ("softbreak", "hardbreak"):
                    paragraph.add_run("\n")
        elif token.type == "fence":
            document.add_paragraph(token.content)
    document.add_heading("会议来源与原文依据", 1)
    referenced = set(re.findall(r"\[(S\d+-P\d+)\]", result["body"]))
    for i, source in enumerate(sources, 1):
        document.add_heading(f"S{i} {source['title']}", 2)
        document.add_paragraph(f"来源：{source['platform']}\n{source.get('url') or '用户提供的文字或文件'}")
        for p in source["paragraphs"]:
            evidence_id = f"S{i}-{p['id']}"
            if evidence_id not in referenced:
                continue
            seconds = int(p["startMs"] / 1000) if p.get("startMs") is not None else None
            locator = (
                f"{seconds // 3600:02}:{seconds // 60 % 60:02}:{seconds % 60:02}" if seconds is not None else p["id"]
            )
            document.add_paragraph(f"[{evidence_id}] {p['speaker']} | {locator}\n{p['text']}")
    for evidence in result.get("formalEvidence", []):
        document.add_paragraph(f"[{evidence['evidence_id']}] {evidence['title']}\n{evidence['excerpt']}")
    if result.get("selectedHistory"):
        document.add_heading("用户选中的历史会议（辅助材料）", 1)
        for meeting in result["selectedHistory"]:
            document.add_heading(f"[{meeting.get('label', '历史会议')}] {meeting['title']}", 2)
            document.add_paragraph(meeting["body"])
    data = io.BytesIO()
    document.save(data)
    return data.getvalue()
